import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import csv
import os

def generate_excel():
    wb = openpyxl.Workbook()
    
    # Define styles
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    pass_font = Font(name="Calibri", size=10, bold=True, color="006100")
    
    fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    fail_font = Font(name="Calibri", size=10, bold=True, color="9C0006")
    
    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    # ----------------------------------------------------
    # Sheet 1: Test Cases
    # ----------------------------------------------------
    ws_tc = wb.active
    ws_tc.title = "Test Cases"
    ws_tc.views.sheetView[0].showGridLines = True
    
    csv_path = os.path.join(os.path.dirname(__file__), "03-physical-product", "TEST_CASES.csv")
    with open(csv_path, mode='r', encoding='utf-8') as f:
        reader = csv.reader(f)
        for r_idx, row in enumerate(reader, 1):
            ws_tc.append(row)
            for c_idx in range(1, len(row) + 1):
                cell = ws_tc.cell(row=r_idx, column=c_idx)
                cell.border = thin_border
                
                if r_idx == 1:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                else:
                    cell.font = regular_font
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    
                    # Highlight Verdict
                    if c_idx == 8:  # Verdict column
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        if cell.value == "PASS":
                            cell.fill = pass_fill
                            cell.font = pass_font
                        elif cell.value == "FAIL":
                            cell.fill = fail_fill
                            cell.font = fail_font

    ws_tc.row_dimensions[1].height = 28

    # Auto-adjust column widths for Test Cases sheet
    for col in ws_tc.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if '\n' in val_str:
                val_str = max(val_str.split('\n'), key=len)
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws_tc.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)

    # ----------------------------------------------------
    # Sheet 2: Test Summary Report
    # ----------------------------------------------------
    ws_sum = wb.create_sheet(title="Test Summary Report")
    ws_sum.views.sheetView[0].showGridLines = True
    
    ws_sum.column_dimensions['A'].width = 6
    ws_sum.column_dimensions['B'].width = 38
    ws_sum.column_dimensions['C'].width = 22
    ws_sum.column_dimensions['D'].width = 45
    
    ws_sum['B2'] = "TEST EXECUTION SUMMARY REPORT"
    ws_sum['B2'].font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    
    metrics = [
        ("Project / SUT", "Bếp hồng ngoại Sunhouse SHD6012A"),
        ("Brand / Model", "SUNHOUSE / SHD6012A (2024)"),
        ("Student Name", "Trần Vũ Quang"),
        ("Student ID", "23120346"),
        ("Course / Assignment", "Software Testing - HW01-AI"),
        ("Execution Date", "04/10/2026"),
        ("Total Test Cases Designed", 15),
        ("Test Cases Executed Live on Hardware", 5),
        ("Dry-Run / Lab Boundary Tested", 10),
        ("Total Passed", 12),
        ("Total Failed", 3),
        ("Pass Rate (%)", "80.0%"),
        ("Fail Rate (%)", "20.0%"),
        ("Total Physical Defects Discovered", 5),
        ("High Severity Defects", 2),
        ("Medium Severity Defects", 3),
        ("AI Missed Edge Cases Identified", 3)
    ]
    
    start_row = 4
    for idx, (label, val) in enumerate(metrics):
        r = start_row + idx
        ws_sum.cell(row=r, column=2, value=label).font = bold_font
        ws_sum.cell(row=r, column=2).border = thin_border
        
        val_cell = ws_sum.cell(row=r, column=3, value=val)
        val_cell.font = regular_font
        val_cell.border = thin_border
        val_cell.alignment = Alignment(horizontal="center")
        
        if label == "Total Passed":
            val_cell.fill = pass_fill
            val_cell.font = pass_font
        elif label in ["Total Failed", "High Severity Defects"]:
            val_cell.fill = fail_fill
            val_cell.font = fail_font

    output_path = os.path.join(os.path.dirname(__file__), "Test_Cases_and_Checklist.xlsx")
    wb.save(output_path)
    print(f"File created successfully: {output_path}")

if __name__ == "__main__":
    generate_excel()
